"""
Fulfillment documents for the shipping role, built from LeafTrade order details.

  - pick list PDF   -> one page per order, modelled on the LeafTrade packing slip
  - summary XLSX    -> every line item for a delivery date, sorted by product and
                       highlighted per pull number, with total units per pull number

Prices and other finance fields never leave this module: only the fields below are read.
LeafTrade's `quantity` is the case count and `units` is the total unit count.
"""

import logging
from io import BytesIO
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Paragraph, Table, TableStyle, Spacer, PageBreak, KeepTogether

from api.leaftrade import get_order_details

logger = logging.getLogger('utils.pick_lists')

_TITLE = ParagraphStyle('title', fontName='Helvetica-Bold', fontSize=14, leading=17)
_CUSTOMER = ParagraphStyle('customer', fontName='Helvetica-Bold', fontSize=11, leading=14)
_LABEL = ParagraphStyle('label', fontName='Helvetica-Bold', fontSize=7, leading=9, textColor=colors.HexColor('#555555'))
_VALUE = ParagraphStyle('value', fontName='Helvetica', fontSize=9, leading=11)

GRID = colors.HexColor('#BBBBBB')

# Cycled per pull number in sort order, so neighbouring groups always differ.
PULL_FILLS = ['FFFF99', 'FFCC99', 'C6EFCE', 'BDD7EE', 'E4C1F9', 'FFC7CE', 'B4E5E5', 'D9D9D9']

SUMMARY_HEADERS = ['customer', 'customer_location', 'product_name', 'pull_number',
                   'quantity', 'cases', 'Total', 'Need Labeled']


def _order(order_id):
    """Raw LeafTrade order dict. Raises when LeafTrade fails - no partial documents."""
    details = get_order_details(str(order_id))
    if not details:
        raise Exception(f"Failed to fetch LeafTrade order {order_id}")
    return details['order']


def _text(value):
    return escape(str(value) if value is not None else '')


def _address(address):
    address = address or {}
    lines = [address.get('street_address_1'), address.get('street_address_2'),
             f"{address.get('city', '')}, {address.get('state', '')} {address.get('postal_code', '')}"]
    return '<br/>'.join(_text(line) for line in lines if line)


def _unit_size(item):
    if item.get('size') is None:
        return ''
    return f"{item['size']:g}{(item.get('size_unit') or '').upper()}"


def _by_product(items):
    return sorted(items, key=lambda item: (item.get('product_name') or '').lower())


def _pick_list(order):
    """Flowables for one order's pick list."""
    location = order.get('dispensary_location') or {}
    customer = f"{location.get('dispensary', {}).get('name', '')} - {location.get('name', '')}"

    header = Table([
        [Paragraph('SHIP TO', _LABEL), Paragraph('DELIVERY DATE', _LABEL)],
        [Paragraph(_address(order.get('shipping_address')), _VALUE), Paragraph(_text(order.get('delivery_date')), _VALUE)],
    ], colWidths=[4.5 * inch, 3 * inch])
    header.setStyle(TableStyle([
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('LEFTPADDING', (0, 0), (-1, -1), 0),
    ]))

    rows = [[Paragraph(label, _LABEL) for label in ('PRODUCT', 'BATCH #', 'PACKAGE ID', 'UNIT SIZE', 'CASES', 'UNITS')]]
    for item in _by_product(order.get('items', [])):
        rows.append([
            Paragraph(_text(item.get('product_name')), _VALUE),
            Paragraph(_text(item.get('batch_number_override') or item.get('batch_ref')), _VALUE),
            Paragraph(_text(item.get('pull_number')), _VALUE),
            Paragraph(_text(_unit_size(item)), _VALUE),
            Paragraph(_text(item.get('quantity')), _VALUE),
            Paragraph(_text(item.get('units')), _VALUE),
        ])
    items = Table(rows, colWidths=[3.2 * inch, 1.5 * inch, 1.1 * inch, 0.7 * inch, 0.5 * inch, 0.5 * inch], repeatRows=1)
    items.setStyle(TableStyle([
        ('LINEBELOW', (0, 0), (-1, -1), 0.5, GRID),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ('LEFTPADDING', (0, 0), (-1, -1), 2),
    ]))

    signatures = Table(
        [[Paragraph(label, _LABEL), Paragraph('DATE', _LABEL)]
         for label in ('ORDER LABELED BY', 'ORDER PACKAGED BY', 'DISPENSARY RECEIVING SIGNATURE')],
        colWidths=[3 * inch, 1.5 * inch], rowHeights=0.45 * inch,
    )
    signatures.setStyle(TableStyle([
        ('LINEABOVE', (0, 0), (0, -1), 0.75, colors.black),
        ('LINEABOVE', (1, 0), (1, -1), 0.75, colors.black),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('LEFTPADDING', (0, 0), (-1, -1), 0),
        ('RIGHTPADDING', (0, 0), (0, -1), 24),
    ]))

    return [
        Paragraph(f"Packing Slip #{_text(order.get('invoice_id'))}", _TITLE),
        Spacer(1, 6),
        Paragraph(_text(customer), _CUSTOMER),
        Paragraph(f"License # {_text(location.get('license_number'))}", _VALUE),
        Spacer(1, 8),
        header,
        Spacer(1, 12),
        items,
        Spacer(1, 36),
        KeepTogether(signatures),
    ]


def build_pick_list_pdf(order_ids):
    """Render pick lists for the given LeafTrade order IDs, one order per page."""
    story = []
    for order_id in order_ids:
        if story:
            story.append(PageBreak())
        story.extend(_pick_list(_order(order_id)))

    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, leftMargin=0.5 * inch, rightMargin=0.5 * inch,
                            topMargin=0.5 * inch, bottomMargin=0.5 * inch, title='Pick Lists')
    doc.build(story)

    logger.info(f"Generated pick list PDF for {len(order_ids)} order(s)")
    return buffer.getvalue()


def build_summary_xlsx(order_ids, delivery_date):
    """Every line item for the orders, alphabetized by product and highlighted per pull number."""
    rows = []
    for order_id in order_ids:
        order = _order(order_id)
        location = order.get('dispensary_location') or {}
        for item in order.get('items', []):
            rows.append([
                location.get('dispensary', {}).get('name', ''),
                location.get('name', ''),
                item.get('product_name') or '',
                item.get('pull_number') or '',
                item.get('units') or 0,
                item.get('quantity') or 0,
            ])
    rows.sort(key=lambda row: (row[2].lower(), row[3], row[0].lower(), row[1].lower()))

    wb = Workbook()
    ws = wb.active
    ws.title = delivery_date
    ws.append(SUMMARY_HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    color_index = -1
    for i, row in enumerate(rows):
        if i == 0 or row[3] != rows[i - 1][3]:
            color_index += 1
            group_units = 0
        group_units += row[4]
        # Total units for the pull number sits on the last row of its group.
        is_last = i == len(rows) - 1 or rows[i + 1][3] != row[3]
        ws.append(row + [group_units if is_last else None, None])
        fill = PatternFill('solid', fgColor=PULL_FILLS[color_index % len(PULL_FILLS)])
        for cell in ws[ws.max_row]:
            cell.fill = fill

    for column, width in zip('ABCDEFGH', [22, 20, 80, 14, 10, 8, 10, 14]):
        ws.column_dimensions[column].width = width
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    buffer = BytesIO()
    wb.save(buffer)

    logger.info(f"Generated fulfillment summary for {delivery_date} ({len(order_ids)} orders, {len(rows)} rows)")
    return buffer.getvalue()
